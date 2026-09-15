# -*- coding: utf-8 -*-
"""
SINAX Merge Service Automated Tests
Tests PDF merging, Word merging, Excel merging, Image conversions, and ZIP bundling.
"""

import os
import shutil
import tempfile
import unittest
from pathlib import Path

from app.services.merge_service import MergeService

class TestMergeService(unittest.TestCase):
    def setUp(self):
        self.test_dir = Path(tempfile.mkdtemp(prefix="sinax_merge_test_"))

    def tearDown(self):
        if self.test_dir.exists():
            shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_detect_merge_mode(self):
        # All PDFs
        pdfs = [self.test_dir / "a.pdf", self.test_dir / "b.pdf"]
        self.assertEqual(MergeService.detect_merge_mode(pdfs)["mode"], "pdf")

        # All DOCX
        docs = [self.test_dir / "a.docx", self.test_dir / "b.docx"]
        self.assertEqual(MergeService.detect_merge_mode(docs)["mode"], "docx")

        # All Excel
        xls = [self.test_dir / "a.xlsx", self.test_dir / "b.xlsx"]
        self.assertEqual(MergeService.detect_merge_mode(xls)["mode"], "excel")

        # All Images
        imgs = [self.test_dir / "a.png", self.test_dir / "b.jpg"]
        self.assertEqual(MergeService.detect_merge_mode(imgs)["mode"], "images")

        # All Videos
        vids = [self.test_dir / "a.mp4", self.test_dir / "b.mkv"]
        self.assertEqual(MergeService.detect_merge_mode(vids)["mode"], "videos")

        # Mixed
        mixed = [self.test_dir / "a.pdf", self.test_dir / "b.png"]
        self.assertEqual(MergeService.detect_merge_mode(mixed)["mode"], "zip")

    def test_pdf_merge(self):
        from pypdf import PdfWriter

        # Create two 1-page PDFs
        p1 = self.test_dir / "doc1.pdf"
        p2 = self.test_dir / "doc2.pdf"

        w1 = PdfWriter()
        w1.add_blank_page(width=200, height=200)
        with open(p1, "wb") as f:
            w1.write(f)

        w2 = PdfWriter()
        w2.add_blank_page(width=200, height=200)
        with open(p2, "wb") as f:
            w2.write(f)

        out_pdf = self.test_dir / "merged.pdf"
        success = MergeService.merge_pdfs([p1, p2], out_pdf)
        self.assertTrue(success)
        self.assertTrue(out_pdf.exists())

        from pypdf import PdfReader
        r = PdfReader(str(out_pdf))
        self.assertEqual(len(r.pages), 2)

    def test_docx_merge(self):
        from docx import Document

        d1_path = self.test_dir / "part1.docx"
        d2_path = self.test_dir / "part2.docx"

        d1 = Document()
        d1.add_heading("Chapter 1", level=1)
        d1.add_paragraph("Content of chapter 1")
        d1.save(str(d1_path))

        d2 = Document()
        d2.add_heading("Chapter 2", level=1)
        d2.add_paragraph("Content of chapter 2")
        d2.save(str(d2_path))

        out_docx = self.test_dir / "book_merged.docx"
        success = MergeService.merge_docx([d1_path, d2_path], out_docx, add_page_breaks=True)
        self.assertTrue(success)
        self.assertTrue(out_docx.exists())

        merged_doc = Document(str(out_docx))
        headings = [p.text for p in merged_doc.paragraphs if p.text in ["Chapter 1", "Chapter 2"]]
        self.assertIn("Chapter 1", headings)
        self.assertIn("Chapter 2", headings)

    def test_excel_sheets_merge(self):
        import openpyxl

        wb1_path = self.test_dir / "data1.xlsx"
        wb2_path = self.test_dir / "data2.xlsx"

        wb1 = openpyxl.Workbook()
        ws1 = wb1.active
        ws1.title = "Sales"
        ws1.append(["Item", "Price"])
        ws1.append(["Pen", 5])
        wb1.save(str(wb1_path))

        wb2 = openpyxl.Workbook()
        ws2 = wb2.active
        ws2.title = "Inventory"
        ws2.append(["Item", "Stock"])
        ws2.append(["Pen", 100])
        wb2.save(str(wb2_path))

        out_xlsx = self.test_dir / "combined.xlsx"
        success = MergeService.merge_excel_sheets([wb1_path, wb2_path], out_xlsx, mode="sheets")
        self.assertTrue(success)
        self.assertTrue(out_xlsx.exists())

        res_wb = openpyxl.load_workbook(str(out_xlsx))
        self.assertEqual(len(res_wb.sheetnames), 2)

    def test_images_to_pdf(self):
        from PIL import Image

        im1_path = self.test_dir / "img1.png"
        im2_path = self.test_dir / "img2.jpg"

        Image.new("RGB", (100, 100), color="red").save(str(im1_path))
        Image.new("RGB", (150, 150), color="blue").save(str(im2_path))

        out_pdf = self.test_dir / "album.pdf"
        success = MergeService.images_to_pdf([im1_path, im2_path], out_pdf)
        self.assertTrue(success)
        self.assertTrue(out_pdf.exists())
        self.assertGreater(out_pdf.stat().st_size, 0)

    def test_stitch_images(self):
        from PIL import Image

        im1_path = self.test_dir / "s1.png"
        im2_path = self.test_dir / "s2.png"

        Image.new("RGB", (100, 50), color="white").save(str(im1_path))
        Image.new("RGB", (100, 50), color="black").save(str(im2_path))

        out_stitch = self.test_dir / "stitched.png"
        success = MergeService.stitch_images([im1_path, im2_path], out_stitch, direction="vertical")
        self.assertTrue(success)
        self.assertTrue(out_stitch.exists())

        res_im = Image.open(str(out_stitch))
        self.assertEqual(res_im.size, (100, 100))

    def test_zip_archive(self):
        f1 = self.test_dir / "note.txt"
        f2 = self.test_dir / "data.csv"
        f1.write_text("hello", encoding="utf-8")
        f2.write_text("a,b,c", encoding="utf-8")

        out_zip = self.test_dir / "bundle.zip"
        success = MergeService.create_zip_archive([f1, f2], out_zip)
        self.assertTrue(success)
        self.assertTrue(out_zip.exists())

        import zipfile
        with zipfile.ZipFile(str(out_zip), "r") as z:
            names = z.namelist()
            self.assertIn("note.txt", names)
            self.assertIn("data.csv", names)

if __name__ == "__main__":
    unittest.main()
