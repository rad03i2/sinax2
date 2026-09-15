# -*- coding: utf-8 -*-
"""
SINAX Merge Service
High-performance, fault-tolerant merging and bundling engine for:
- PDF documents (pypdf)
- Word documents (python-docx)
- Excel workbooks (openpyxl)
- Images to PDF and image stitching (Pillow)
- Heterogeneous and multimedia files into ZIP archives (zipfile)
"""

import os
import zipfile
from pathlib import Path
from typing import List, Optional, Callable, Dict, Any

from app.core.logger import get_logger

logger = get_logger("merge_service")


class MergeService:
    """Core business logic for file merging and archiving."""

    @staticmethod
    def detect_merge_mode(file_paths: List[Path]) -> Dict[str, Any]:
        """
        Analyzes a list of file paths and automatically determines the most appropriate
        merge mode, along with available options and suggested output filename.
        """
        if not file_paths:
            return {
                "mode": "zip",
                "label_ar": "أرشيف مضغوط ZIP",
                "extension": ".zip",
                "options": [],
                "can_merge_directly": False
            }

        exts = [p.suffix.lower() for p in file_paths]
        all_pdf = all(e == ".pdf" for e in exts)
        all_docx = all(e in [".docx", ".doc"] for e in exts)
        all_excel = all(e in [".xlsx", ".xls"] for e in exts)
        image_exts = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".tiff", ".gif"}
        all_images = all(e in image_exts for e in exts)
        video_exts = {".mp4", ".mkv", ".avi", ".mov", ".wmv", ".flv", ".webm"}
        all_videos = all(e in video_exts for e in exts)

        first_stem = file_paths[0].stem

        if all_pdf:
            return {
                "mode": "pdf",
                "label_ar": "دمج ملفات PDF في مستند واحد",
                "extension": ".pdf",
                "suggested_name": f"{first_stem}_مدمج.pdf",
                "options": ["دمج قياسي مع الحفاظ على الصفحات والعلامات"],
                "can_merge_directly": True
            }
        elif all_docx:
            return {
                "mode": "docx",
                "label_ar": "دمج مستندات Word في مستند واحد",
                "extension": ".docx",
                "suggested_name": f"{first_stem}_مدمج.docx",
                "options": ["فواصل صفحات بين المستندات", "الحفاظ على التنسيق والجداول"],
                "can_merge_directly": True
            }
        elif all_excel:
            return {
                "mode": "excel",
                "label_ar": "دمج جداول Excel",
                "extension": ".xlsx",
                "suggested_name": f"{first_stem}_مجمّع.xlsx",
                "options": ["كل ملف في ورقة عمل مستقلة (Sheet)", "دمج الصفوف المتوافقة في ورقة واحدة"],
                "can_merge_directly": True
            }
        elif all_images:
            return {
                "mode": "images",
                "label_ar": "تجميع ودمج الصور",
                "extension": ".pdf",
                "suggested_name": f"ألبوم_صور_مدمج.pdf",
                "options": ["تحويل الصور إلى مستند PDF واحد", "دمج الصور رأسياً (Vertical Stitch)", "دمج الصور أفقياً (Horizontal Stitch)", "حفظ الصور داخل أرشيف ZIP"],
                "can_merge_directly": True
            }
        elif all_videos:
            return {
                "mode": "videos",
                "label_ar": "تجميع ملفات الفيديو",
                "extension": ".zip",
                "suggested_name": f"فيديوهات_مجمعة.zip",
                "options": ["تجميع في أرشيف مضغوط ZIP (موصى به)", "دمج بواسطة FFmpeg (إن وُجد)"],
                "can_merge_directly": False
            }
        else:
            return {
                "mode": "zip",
                "label_ar": "حزم الملفات المتنوعة في أرشيف مضغوط ZIP",
                "extension": ".zip",
                "suggested_name": "ملفات_مجمعة.zip",
                "options": ["أرشيف ZIP مضغوط عالي الكفاءة"],
                "can_merge_directly": False
            }

    @staticmethod
    def merge_pdfs(
        input_paths: List[Path],
        output_path: Path,
        progress_callback: Optional[Callable[[int, int, str], None]] = None
    ) -> bool:
        """Merges multiple PDF files into one output file using pypdf."""
        try:
            from pypdf import PdfWriter
        except ImportError:
            logger.error("pypdf library not installed.")
            raise RuntimeError("حزمة pypdf غير مثبتة في النظام.")

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        writer = PdfWriter()
        total = len(input_paths)

        try:
            for idx, p in enumerate(input_paths):
                if progress_callback:
                    progress_callback(idx + 1, total, p.name)
                writer.append(str(p))

            with open(output_path, "wb") as f_out:
                writer.write(f_out)
            writer.close()
            logger.info(f"Successfully merged {total} PDFs into {output_path}")
            return True
        except Exception as e:
            writer.close()
            logger.error(f"Error during PDF merge: {e}")
            raise RuntimeError(f"فشل دمج ملفات PDF: {str(e)}")

    @staticmethod
    def merge_docx(
        input_paths: List[Path],
        output_path: Path,
        add_page_breaks: bool = True,
        progress_callback: Optional[Callable[[int, int, str], None]] = None
    ) -> bool:
        """Merges multiple DOCX documents into one document using python-docx."""
        try:
            from docx import Document
        except ImportError:
            logger.error("python-docx library not installed.")
            raise RuntimeError("حزمة python-docx غير مثبتة في النظام.")

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        if not input_paths:
            return False

        total = len(input_paths)
        try:
            master = Document(str(input_paths[0]))
            if progress_callback:
                progress_callback(1, total, input_paths[0].name)

            for idx, path in enumerate(input_paths[1:], start=2):
                if progress_callback:
                    progress_callback(idx, total, path.name)

                if add_page_breaks:
                    master.add_page_break()

                sub_doc = Document(str(path))
                for element in sub_doc.element.body:
                    # Skip section properties on intermediary docs to keep master geometry
                    if element.tag.endswith('sectPr'):
                        continue
                    master.element.body.append(element)

            master.save(str(output_path))
            logger.info(f"Successfully merged {total} DOCX files into {output_path}")
            return True
        except Exception as e:
            logger.error(f"Error during Word merge: {e}")
            raise RuntimeError(f"فشل دمج مستندات Word: {str(e)}")

    @staticmethod
    def merge_excel_sheets(
        input_paths: List[Path],
        output_path: Path,
        mode: str = "sheets",  # 'sheets' or 'concat'
        progress_callback: Optional[Callable[[int, int, str], None]] = None
    ) -> bool:
        """Merges multiple Excel files into one workbook using openpyxl."""
        try:
            import openpyxl
            from openpyxl import Workbook
        except ImportError:
            logger.error("openpyxl library not installed.")
            raise RuntimeError("حزمة openpyxl غير مثبتة في النظام.")

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        total = len(input_paths)
        master_wb = Workbook()
        # Remove the initial default sheet
        default_sheet = master_wb.active

        try:
            if mode == "sheets":
                used_sheet_names = set()
                for idx, path in enumerate(input_paths):
                    if progress_callback:
                        progress_callback(idx + 1, total, path.name)

                    wb = openpyxl.load_workbook(str(path), data_only=True)
                    file_stem = path.stem[:18]  # Limit prefix

                    for s_name in wb.sheetnames:
                        src_sheet = wb[s_name]
                        # Create unique valid sheet name (Excel max 31 chars)
                        clean_name = f"{file_stem}_{s_name}"[:31]
                        counter = 1
                        while clean_name in used_sheet_names:
                            clean_name = f"{clean_name[:27]}_{counter}"
                            counter += 1
                        used_sheet_names.add(clean_name)

                        dest_sheet = master_wb.create_sheet(title=clean_name)
                        for row in src_sheet.iter_rows(values_only=True):
                            dest_sheet.append(list(row))

                if default_sheet in master_wb.worksheets and len(master_wb.worksheets) > 1:
                    master_wb.remove(default_sheet)

            elif mode == "concat":
                # Merge into a single master sheet
                dest_sheet = master_wb.active
                dest_sheet.title = "البيانات المدمجة"
                header_written = False

                for idx, path in enumerate(input_paths):
                    if progress_callback:
                        progress_callback(idx + 1, total, path.name)

                    wb = openpyxl.load_workbook(str(path), data_only=True)
                    first_sheet = wb.active
                    rows = list(first_sheet.iter_rows(values_only=True))
                    if not rows:
                        continue

                    if not header_written:
                        dest_sheet.append(list(rows[0]))
                        header_written = True

                    # Append data rows
                    for r in rows[1:]:
                        dest_sheet.append(list(r))

            master_wb.save(str(output_path))
            logger.info(f"Successfully merged {total} Excel files into {output_path}")
            return True
        except Exception as e:
            logger.error(f"Error during Excel merge: {e}")
            raise RuntimeError(f"فشل دمج جداول Excel: {str(e)}")

    @staticmethod
    def images_to_pdf(
        image_paths: List[Path],
        output_pdf_path: Path,
        progress_callback: Optional[Callable[[int, int, str], None]] = None
    ) -> bool:
        """Combines multiple image files into a single high-quality PDF document."""
        try:
            from PIL import Image
        except ImportError:
            logger.error("Pillow library not installed.")
            raise RuntimeError("حزمة Pillow غير مثبتة في النظام.")

        output_pdf_path = Path(output_pdf_path)
        output_pdf_path.parent.mkdir(parents=True, exist_ok=True)

        if not image_paths:
            return False

        total = len(image_paths)
        converted_images = []

        try:
            for idx, p in enumerate(image_paths):
                if progress_callback:
                    progress_callback(idx + 1, total, p.name)

                img = Image.open(str(p))
                # Convert RGBA/P/LA to RGB for clean PDF output
                if img.mode != "RGB":
                    img = img.convert("RGB")
                converted_images.append(img)

            if converted_images:
                first = converted_images[0]
                rest = converted_images[1:] if len(converted_images) > 1 else []
                first.save(
                    str(output_pdf_path),
                    "PDF",
                    resolution=100.0,
                    save_all=True,
                    append_images=rest
                )

            logger.info(f"Successfully converted {total} images into PDF {output_pdf_path}")
            return True
        except Exception as e:
            logger.error(f"Error converting images to PDF: {e}")
            raise RuntimeError(f"فشل تحويل الصور إلى PDF: {str(e)}")

    @staticmethod
    def stitch_images(
        image_paths: List[Path],
        output_path: Path,
        direction: str = "vertical",  # 'vertical' or 'horizontal'
        progress_callback: Optional[Callable[[int, int, str], None]] = None
    ) -> bool:
        """Stitches multiple images into a single long or panoramic image."""
        try:
            from PIL import Image
        except ImportError:
            logger.error("Pillow library not installed.")
            raise RuntimeError("حزمة Pillow غير مثبتة في النظام.")

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        if not image_paths:
            return False

        total = len(image_paths)
        loaded = []

        try:
            for idx, p in enumerate(image_paths):
                if progress_callback:
                    progress_callback(idx + 1, total, p.name)
                loaded.append(Image.open(str(p)))

            if direction == "vertical":
                max_w = max(im.width for im in loaded)
                total_h = sum(im.height for im in loaded)
                stitched = Image.new("RGBA", (max_w, total_h), (255, 255, 255, 0))
                y_offset = 0
                for im in loaded:
                    # Center align horizontally if widths differ
                    x_offset = (max_w - im.width) // 2
                    stitched.paste(im, (x_offset, y_offset))
                    y_offset += im.height
            else:
                total_w = sum(im.width for im in loaded)
                max_h = max(im.height for im in loaded)
                stitched = Image.new("RGBA", (total_w, max_h), (255, 255, 255, 0))
                x_offset = 0
                for im in loaded:
                    # Center align vertically if heights differ
                    y_offset = (max_h - im.height) // 2
                    stitched.paste(im, (x_offset, y_offset))
                    x_offset += im.width

            # Convert to RGB if saving as JPEG
            if output_path.suffix.lower() in [".jpg", ".jpeg"]:
                stitched = stitched.convert("RGB")

            stitched.save(str(output_path))
            logger.info(f"Successfully stitched {total} images into {output_path}")
            return True
        except Exception as e:
            logger.error(f"Error stitching images: {e}")
            raise RuntimeError(f"فشل دمج الصور: {str(e)}")

    @staticmethod
    def create_zip_archive(
        file_paths: List[Path],
        output_zip: Path,
        progress_callback: Optional[Callable[[int, int, str], None]] = None
    ) -> bool:
        """Packages heterogeneous files into a zip archive with deflate compression."""
        output_zip = Path(output_zip)
        output_zip.parent.mkdir(parents=True, exist_ok=True)

        total = len(file_paths)
        try:
            with zipfile.ZipFile(str(output_zip), 'w', compression=zipfile.ZIP_DEFLATED) as zf:
                for idx, path in enumerate(file_paths):
                    if progress_callback:
                        progress_callback(idx + 1, total, path.name)
                    if path.is_file():
                        zf.write(str(path), arcname=path.name)

            logger.info(f"Successfully created ZIP archive {output_zip} with {total} files")
            return True
        except Exception as e:
            logger.error(f"Error creating ZIP archive: {e}")
            raise RuntimeError(f"فشل إنشاء الأرشيف المضغوط ZIP: {str(e)}")
